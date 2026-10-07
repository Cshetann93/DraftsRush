from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("downloads")
OUT.mkdir(exist_ok=True)
NAVY, GOLD, PALE, MUTED = "233B53", "E3B341", "F2F5F7", "5B6873"

def shade(cell, color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)

def margins(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", 100), ("start", 100), ("bottom", 100), ("end", 100)):
        node = OxmlElement("w:" + side); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); mar.append(node)
    tcPr.append(mar)

def make_doc(title, subtitle, landscape=False):
    d = Document()
    sec = d.sections[0]
    sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55)
    sec.left_margin = Inches(.55); sec.right_margin = Inches(.55)
    if landscape:
        sec.orientation = 1
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    d.styles["Normal"].font.name = "Aptos"
    d.styles["Normal"].font.size = Pt(9)
    d.styles["Title"].font.name = "Aptos Display"
    d.styles["Title"].font.size = Pt(23)
    d.styles["Title"].font.bold = True
    d.styles["Title"].font.color.rgb = RGBColor.from_string(NAVY)
    p=d.add_paragraph(style="Title"); p.add_run(title)
    p=d.add_paragraph(subtitle); p.paragraph_format.space_after=Pt(8)
    for r in p.runs: r.font.color.rgb=RGBColor.from_string(MUTED)
    t=d.add_table(rows=1, cols=1); shade(t.cell(0,0), GOLD); t.cell(0,0).text=""
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=f.add_run("DRAFTSRUSH  •  FREE FILMMAKING TEMPLATE  •  EDITABLE IN WORD")
    r.bold=True; r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string(MUTED)
    return d

def fields(d, labels, cols=2):
    t=d.add_table(rows=(len(labels)+cols-1)//cols, cols=cols); t.style="Table Grid"
    for i,label in enumerate(labels):
        c=t.cell(i//cols,i%cols); c.text=""
        p=c.paragraphs[0]; r=p.add_run(label+"\n"); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); r.font.size=Pt(8)
        p.add_run("________________________________________")
        margins(c)
    d.add_paragraph()

def section(d, title):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(5)
    r=p.add_run(title.upper()); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); r.font.size=Pt(10)

def table(d, headers, rows=(), blanks=0):
    t=d.add_table(rows=1, cols=len(headers)); t.style="Table Grid"
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); margins(c)
        for r in c.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for row in list(rows)+[[""]*len(headers) for _ in range(blanks)]:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=val; margins(cells[i])
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(8)
    d.add_paragraph()

def save(name,d): d.save(OUT/name)

d=make_doc("Script Breakdown Sheet","Use one sheet per scene to track cast, props, wardrobe, locations, and production needs.")
fields(d,["Production / Project","Scene Number","Script Page(s)","INT. / EXT.","Location","Day / Night","Scene Heading","Estimated Duration"])
section(d,"Scene summary"); table(d,["Brief action / scene summary"],[["\n\n"]])
section(d,"Departments & requirements")
table(d,["Category","Items / details","Ready?"],[
["Cast / background","","☐ Yes  ☐ No"],["Props / set dressing","","☐ Yes  ☐ No"],["Wardrobe / makeup","","☐ Yes  ☐ No"],
["Special effects / stunts","","☐ Yes  ☐ No"],["Vehicles / animals","","☐ Yes  ☐ No"],["Sound / music","","☐ Yes  ☐ No"],
["Camera / lighting","","☐ Yes  ☐ No"],["Permissions / safety","","☐ Yes  ☐ No"]],2)
fields(d,["Continuity notes","Owner / department","Open questions / follow-up"],1)
save("script-breakdown-sheet-template.docx",d)

d=make_doc("Shot List","Plan coverage before the shoot. Add a row for each planned setup or take.")
fields(d,["Production / Project","Director","Date","Scene(s)","Location","DOP / Camera"])
table(d,["Shot #","Scene","Shot size / angle","Movement","Action / framing","Lens / notes","Priority","Status"],blanks=12)
section(d,"On-set checklist")
d.add_paragraph("☐ Camera / media checked     ☐ Batteries charged     ☐ Sound checked     ☐ Exposure checked")
d.add_paragraph("☐ Master / establishing shot     ☐ Coverage / reverses     ☐ Inserts / cutaways     ☐ Room tone")
fields(d,["Continuity notes","Missing coverage / pickups"],1)
save("shot-list-template.docx",d)

d=make_doc("Storyboard Template","Sketch or paste a frame into each panel; describe the action and camera notes.",True)
fields(d,["Project","Scene","Director / Artist","Date"],4)
t=d.add_table(rows=0,cols=3); t.style="Table Grid"
for i in range(1,7):
    cs=t.add_row().cells
    cs[0].text=f"FRAME {i}\n\n\n\n\n\n"; shade(cs[0],PALE)
    cs[1].text="ACTION / COMPOSITION\n\n\n"; cs[2].text="CAMERA / SOUND\n\n\n"
    for c in cs: margins(c)
d.add_paragraph(); fields(d,["Sequence notes / transitions","Props / continuity"],1)
save("storyboard-template.docx",d)

d=make_doc("Film Call Sheet","Share with the production team. Confirm all times, locations, and emergency details before distribution.")
fields(d,["Production / Project","Shoot Date","Shoot Day #","Call Time","Estimated Wrap","Weather / Conditions","Producer / AD","Version / Updated"])
section(d,"Locations & contacts"); table(d,["Location","Address / access notes","Parking / basecamp","Contact / phone"],blanks=3)
section(d,"Schedule"); table(d,["Time","Scene(s)","Location","Cast / department","Notes"],blanks=7)
section(d,"Cast call"); table(d,["Cast # / name","Call time","HMU / wardrobe","Scene(s)","Notes"],blanks=5)
section(d,"Crew / important contacts"); table(d,["Role","Name","Phone / radio","Notes"],blanks=5)
section(d,"Safety & reminders")
d.add_paragraph("☐ Emergency exits reviewed     ☐ First-aid contact confirmed     ☐ Weather risks checked")
d.add_paragraph("☐ Stunts / effects briefed     ☐ Permits / releases confirmed     ☐ Transport plan confirmed")
fields(d,["Nearest hospital / emergency address","Emergency contact / safety lead"],1)
save("film-call-sheet-template.docx",d)

d=make_doc("Shooting Schedule","Build a practical day plan with setup, cast, location, and buffer time.")
fields(d,["Production / Project","Shoot Date","Director","1st AD","Location / Unit","Planned Wrap"])
table(d,["Start","End","Duration","Scene","Setup / activity","Location","Cast / department","Notes"],blanks=12)
section(d,"Day plan checks")
d.add_paragraph("☐ Meal / rest breaks planned     ☐ Company moves included     ☐ Lighting / camera resets included")
d.add_paragraph("☐ Travel time checked     ☐ Weather / daylight considered     ☐ Wrap planned")
fields(d,["Contingency / weather plan","Overtime approval / contact","End-of-day notes"],1)
save("shooting-schedule-template.docx",d)
print("Generated five editable DOCX templates in downloads/.")
